"""
EL RINCÓN 3D LAB — ERP Excel Builder
Part 1: Setup, styles, and sheets 01-05
"""
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.chart import BarChart, Reference, LineChart, PieChart
from openpyxl.chart.series import DataPoint
import datetime

# ─────────────────────────────────────────
# COLOR PALETTE (F1 / Dark theme)
# ─────────────────────────────────────────
C_BLACK    = "0D0D0D"
C_DARK     = "1A1A2E"
C_MID      = "16213E"
C_ACCENT   = "E31837"   # F1 Red
C_GOLD     = "FFD700"
C_GREEN    = "00C851"
C_TEAL     = "00BCD4"
C_WHITE    = "FFFFFF"
C_LGRAY    = "F5F5F5"
C_GRAY     = "CCCCCC"
C_DGRAY    = "555555"
C_ORANGE   = "FF6F00"
C_HEADER   = "1A1A2E"
C_SUBHEAD  = "E31837"
C_ROW_ALT  = "F0F4FF"
C_ROW_EVEN = "FFFFFF"
C_BORDER   = "DDDDDD"

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(bold=False, color=C_BLACK, size=10, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic,
                name="Segoe UI")

def align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def border_thin():
    s = Side(style="thin", color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)

def border_med():
    s = Side(style="medium", color=C_DGRAY)
    return Border(left=s, right=s, top=s, bottom=s)

def style_header_row(ws, row, cols, bg=C_HEADER, fg=C_WHITE, size=9):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill(bg)
        cell.font = font(bold=True, color=fg, size=size)
        cell.alignment = align("center", "center")
        cell.border = border_thin()

def style_title(ws, cell_ref, text, bg=C_ACCENT, fg=C_WHITE, size=14):
    cell = ws[cell_ref]
    cell.value = text
    cell.fill = fill(bg)
    cell.font = font(bold=True, color=fg, size=size)
    cell.alignment = align("center", "center")

def add_table(ws, ref, name, style="TableStyleMedium2"):
    tbl = Table(displayName=name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name=style, showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tbl)

def freeze(ws, cell):
    ws.freeze_panes = cell

def set_col_width(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def set_row_height(ws, row, h):
    ws.row_dimensions[row].height = h

def tab_color(ws, hex_color):
    ws.sheet_properties.tabColor = hex_color

# ─────────────────────────────────────────
# WORKBOOK
# ─────────────────────────────────────────
wb = Workbook()
wb.remove(wb.active)  # remove default sheet

# ═══════════════════════════════════════════════════════
# SHEET 01 — CONFIGURACIÓN
# ═══════════════════════════════════════════════════════
ws = wb.create_sheet("01_CONFIGURACION")
tab_color(ws, C_ACCENT)
ws.sheet_view.showGridLines = False

# Title
ws.merge_cells("A1:D1")
style_title(ws, "A1", "⚙️  EL RINCÓN 3D LAB — CONFIGURACIÓN GLOBAL", C_ACCENT, C_WHITE, 14)
set_row_height(ws, 1, 32)

ws.merge_cells("A2:D2")
ws["A2"].value = "Variables globales del sistema · Modificar solo en esta hoja"
ws["A2"].fill = fill(C_DARK)
ws["A2"].font = font(color=C_GRAY, size=9, italic=True)
ws["A2"].alignment = align("center", "center")
set_row_height(ws, 2, 18)

# Section headers + rows
sections = [
    ("TIPO DE CAMBIO", [
        ("TC_USD_UYU", "Tipo de cambio USD/UYU",     42.50,  "$ UYU por 1 USD"),
        ("TC_USD_ARS", "Tipo de cambio USD/ARS",    880.00,  "$ ARS por 1 USD"),
    ]),
    ("COMISIONES Y COSTOS PLATAFORMAS", [
        ("COM_MERCADOPAGO",  "Comisión Mercado Pago (%)",   3.49, "% sobre venta"),
        ("COM_WOOCOMMERCE",  "Comisión WooCommerce (%)",    2.00, "% sobre venta"),
        ("COM_MERCADOLIBRE", "Comisión Mercado Libre (%)", 13.00, "% sobre venta"),
        ("COSTO_DAC",        "Costo DAC promedio (UYU)",   45.00, "Por envío"),
    ]),
    ("PUBLICIDAD", [
        ("PRESUP_META_MES",   "Presupuesto Meta Ads mensual",  5000.00, "UYU/mes"),
        ("PRESUP_GOOGLE_MES", "Presupuesto Google Ads mensual", 0.00,   "UYU/mes"),
    ]),
    ("OBJETIVOS", [
        ("OBJ_VENTAS_MES",    "Objetivo ventas mensual (UYU)",  80000.00, "UYU"),
        ("OBJ_UTILIDAD_MES",  "Objetivo utilidad mensual (UYU)", 25000.00, "UYU"),
        ("MARGEN_OBJETIVO",   "Margen objetivo (%)",                35.00, "%"),
        ("TICKET_OBJETIVO",   "Ticket promedio objetivo (UYU)",  1200.00, "UYU"),
    ]),
    ("COSTOS FIJOS MENSUALES", [
        ("CF_ELECTRICIDAD",  "Electricidad",          1200.00, "UYU/mes"),
        ("CF_INTERNET",      "Internet",               800.00, "UYU/mes"),
        ("CF_SOFTWARE",      "Software / SaaS",        500.00, "UYU/mes"),
        ("CF_FILAMENTO",     "Filamento base mensual",3000.00, "UYU/mes"),
        ("CF_OTROS",         "Otros costos fijos",     500.00, "UYU/mes"),
    ]),
    ("IMPRESIÓN 3D", [
        ("COSTO_KWH",        "Costo kWh (UYU)",         6.50, "UYU"),
        ("CONSUMO_HORA",     "Consumo impresora (kWh/h)", 0.25, "kWh por hora"),
        ("COSTO_FILAMENTO_KG","Costo filamento (UYU/kg)", 600.00, "UYU por kg"),
        ("GRAMOS_POR_HORA",  "Gramos filamento / hora",   30.00, "g/h promedio"),
    ]),
]

row = 4
for section_title, variables in sections:
    # Section header
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"].value = section_title
    ws[f"A{row}"].fill = fill(C_MID)
    ws[f"A{row}"].font = font(bold=True, color=C_GOLD, size=9)
    ws[f"A{row}"].alignment = align("left", "center")
    ws[f"A{row}"].border = border_thin()
    set_row_height(ws, row, 20)
    row += 1

    for var_name, label, default_val, note in variables:
        ws[f"A{row}"].value = var_name
        ws[f"A{row}"].fill = fill(C_DARK)
        ws[f"A{row}"].font = font(bold=True, color=C_TEAL, size=9)
        ws[f"A{row}"].alignment = align("left", "center")
        ws[f"A{row}"].border = border_thin()

        ws[f"B{row}"].value = label
        ws[f"B{row}"].fill = fill(C_LGRAY)
        ws[f"B{row}"].font = font(size=9)
        ws[f"B{row}"].alignment = align("left", "center")
        ws[f"B{row}"].border = border_thin()

        ws[f"C{row}"].value = default_val
        ws[f"C{row}"].fill = fill("FFFDE7")
        ws[f"C{row}"].font = font(bold=True, size=10, color=C_DARK)
        ws[f"C{row}"].alignment = align("right", "center")
        ws[f"C{row}"].border = border_med()
        ws[f"C{row}"].number_format = '#,##0.00'

        ws[f"D{row}"].value = note
        ws[f"D{row}"].fill = fill(C_LGRAY)
        ws[f"D{row}"].font = font(italic=True, size=8, color=C_DGRAY)
        ws[f"D{row}"].alignment = align("left", "center")
        ws[f"D{row}"].border = border_thin()

        set_row_height(ws, row, 18)
        row += 1

    row += 1  # blank between sections

# Column widths
set_col_width(ws, {"A": 22, "B": 38, "C": 18, "D": 30})

# ═══════════════════════════════════════════════════════
# SHEET 02 — PRODUCTOS
# ═══════════════════════════════════════════════════════
ws2 = wb.create_sheet("02_PRODUCTOS")
tab_color(ws2, "1565C0")
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:R1")
style_title(ws2, "A1", "📦  PRODUCTOS — BASE MAESTRA", "1565C0", C_WHITE, 13)
set_row_height(ws2, 1, 28)

headers_prod = [
    ("ID",              8),
    ("SKU",             14),
    ("Nombre",          30),
    ("Categoría",       18),
    ("Subcategoría",    18),
    ("Proveedor",       18),
    ("Costo MP",        12),
    ("Costo Imp.",      12),
    ("Costo DTF",       12),
    ("Costo Pack.",     12),
    ("Costo Envío",     12),
    ("Costo Total",     13),
    ("Precio Venta",    13),
    ("Ganancia",        12),
    ("Margen %",        10),
    ("Peso (g)",        10),
    ("Tiempo Imp.(h)",  14),
    ("Activo",           9),
]

for col_idx, (h, w) in enumerate(headers_prod, 1):
    cell = ws2.cell(row=2, column=col_idx, value=h)
    cell.fill = fill("1565C0")
    cell.font = font(bold=True, color=C_WHITE, size=9)
    cell.alignment = align("center", "center")
    cell.border = border_thin()
    ws2.column_dimensions[get_column_letter(col_idx)].width = w
set_row_height(ws2, 2, 22)

# Sample data rows
categorias = ["IMPRESIÓN 3D", "INDUMENTARIA", "TERMOS Y BEBIDAS", "ILUMINACIÓN", "PERSONALIZADOS"]
sample_products = [
    (1,"IMP3D-001","Alerón F1 Escala 1:18","IMPRESIÓN 3D","Alerones F1","Filamento PLA",85,45,0,20,80,230,450,220,48.9,120,4.5,"SI"),
    (2,"IMP3D-002","Circuito Monaco Mini","IMPRESIÓN 3D","Circuitos","Filamento PLA",65,35,0,15,80,195,380,185,48.7,90,3.5,"SI"),
    (3,"IMP3D-003","Mini Garaje F1","IMPRESIÓN 3D","Mini Garajes","Filamento PLA",120,60,0,25,80,285,550,265,48.2,180,6.0,"SI"),
    (4,"IMP3D-004","Llavero F1","IMPRESIÓN 3D","Llaveros","Filamento PLA",15,8,0,8,0,31,90,59,65.6,20,0.8,"SI"),
    (5,"IMP3D-005","Imán F1","IMPRESIÓN 3D","Imanes","Filamento PLA",12,6,0,5,0,23,70,47,67.1,15,0.5,"SI"),
    (6,"IMP3D-006","Portallaves F1","IMPRESIÓN 3D","Portallaves","Filamento PLA",35,18,0,12,0,65,180,115,63.9,50,1.5,"SI"),
    (7,"IMP3D-007","Calendario F1 2025","IMPRESIÓN 3D","Calendarios","Filamento PLA",80,40,0,20,0,140,320,180,56.3,110,3.5,"SI"),
    (8,"INDUM-001","Hoodie F1","INDUMENTARIA","Hoodies","Proveedor Textil",800,0,350,40,120,1310,2200,890,40.5,0,0,"SI"),
    (9,"INDUM-002","Remera F1","INDUMENTARIA","Remeras","Proveedor Textil",350,0,180,25,80,635,1100,465,42.3,0,0,"SI"),
    (10,"INDUM-003","Gorra F1","INDUMENTARIA","Gorras","Proveedor Textil",420,0,120,20,80,640,1200,560,46.7,0,0,"SI"),
    (11,"TERM-001","Mate F1 Personalizado","TERMOS Y BEBIDAS","Mates","Proveedor Mates",450,0,200,30,80,760,1400,640,45.7,0,0,"SI"),
    (12,"TERM-002","Vaso Térmico F1","TERMOS Y BEBIDAS","Vasos Térmicos","Proveedor Mates",520,0,180,35,80,815,1500,685,45.7,0,0,"SI"),
    (13,"ILUM-001","Lámpara LED F1","ILUMINACIÓN","Lámparas LED","Proveedor LED",600,80,0,45,120,845,1800,955,53.1,0,0,"SI"),
    (14,"ILUM-002","Neón LED F1","ILUMINACIÓN","Neones LED","Proveedor LED",1200,0,0,60,150,1410,3200,1790,55.9,0,0,"SI"),
    (15,"PERS-001","Producto Personalizado","PERSONALIZADOS","A medida","Variable",0,0,0,0,0,0,0,0,0,0,0,"SI"),
]

for r_idx, row_data in enumerate(sample_products, 3):
    bg = C_ROW_ALT if r_idx % 2 == 0 else C_ROW_EVEN
    for c_idx, val in enumerate(row_data, 1):
        cell = ws2.cell(row=r_idx, column=c_idx, value=val)
        cell.border = border_thin()
        cell.alignment = align("center" if c_idx in [1,4,5,6,18] else "right" if c_idx >= 7 else "left", "center")
        if c_idx == 3:
            cell.alignment = align("left", "center")
        if c_idx in [7,8,9,10,11,12,13,14]:
            cell.number_format = '#,##0'
            cell.fill = fill(bg)
        elif c_idx == 15:
            cell.number_format = '0.0"%"'
            if isinstance(val, (int, float)) and val > 0:
                if val >= 50:
                    cell.fill = fill("C8E6C9")
                elif val >= 35:
                    cell.fill = fill("FFF9C4")
                else:
                    cell.fill = fill("FFCDD2")
            else:
                cell.fill = fill(bg)
        elif c_idx == 18:
            cell.fill = fill("C8E6C9") if val == "SI" else fill("FFCDD2")
            cell.font = font(bold=True, size=9, color="1B5E20" if val == "SI" else "B71C1C")
            cell.alignment = align("center", "center")
        else:
            cell.fill = fill(bg)
            cell.font = font(size=9)
    set_row_height(ws2, r_idx, 16)

# Formulas for row 3 (Costo Total, Ganancia, Margen) — auto-calc
# Col L = Costo Total = G+H+I+J+K
# Col N = Ganancia = M - L
# Col O = Margen = N/M
for r in range(3, 3 + len(sample_products)):
    ws2.cell(row=r, column=12).value = f"=G{r}+H{r}+I{r}+J{r}+K{r}"
    ws2.cell(row=r, column=14).value = f"=M{r}-L{r}"
    ws2.cell(row=r, column=15).value = f"=IF(M{r}>0,N{r}/M{r}*100,0)"
    ws2.cell(row=r, column=15).number_format = '0.0'

# Data validation — Categoría
dv_cat = DataValidation(
    type="list",
    formula1='"IMPRESIÓN 3D,INDUMENTARIA,TERMOS Y BEBIDAS,ILUMINACIÓN,PERSONALIZADOS"',
    allow_blank=True, showErrorMessage=True,
    errorTitle="Categoría inválida", error="Selecciona una categoría válida"
)
ws2.add_data_validation(dv_cat)
dv_cat.add(f"D3:D10000")

dv_activo = DataValidation(type="list", formula1='"SI,NO"', allow_blank=False)
ws2.add_data_validation(dv_activo)
dv_activo.add("R3:R10000")

add_table(ws2, f"A2:R{2+len(sample_products)}", "TBL_Productos", "TableStyleMedium2")
freeze(ws2, "D3")

print("Sheets 01-02 done")

# ═══════════════════════════════════════════════════════
# SHEET 03 — CLIENTES (CRM)
# ═══════════════════════════════════════════════════════
ws3 = wb.create_sheet("03_CLIENTES")
tab_color(ws3, "2E7D32")
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:T1")
style_title(ws3, "A1", "👤  CLIENTES — CRM MAESTRO", "2E7D32", C_WHITE, 13)
set_row_height(ws3, 1, 28)

headers_cli = [
    ("ID Cliente",8),("Nombre",15),("Apellido",15),("Teléfono",14),
    ("Email",22),("Instagram",16),("Ciudad",14),("Departamento",14),("País",10),
    ("1ra Compra",13),("Últ. Compra",13),("Cant. Pedidos",12),
    ("Facturación",13),("Ganancia",13),("Ticket Prom.",13),
    ("Canal Adq.",14),("Segmento",12),("R",6),("F",6),("M",6),
]

for c_idx, (h, w) in enumerate(headers_cli, 1):
    cell = ws3.cell(row=2, column=c_idx, value=h)
    cell.fill = fill("2E7D32")
    cell.font = font(bold=True, color=C_WHITE, size=9)
    cell.alignment = align("center", "center")
    cell.border = border_thin()
    ws3.column_dimensions[get_column_letter(c_idx)].width = w
set_row_height(ws3, 2, 22)

sample_clients = [
    (1,"Martín","García","099123456","martin@email.com","@martin_f1","Montevideo","Montevideo","UY","2024-01-15","2025-05-10",8,12500,4800,1562,"Instagram","VIP",5,4,5),
    (2,"Laura","Rodríguez","098234567","laura@email.com","@laura_gr","Salto","Salto","UY","2024-03-20","2025-04-28",3,3800,1420,1267,"WhatsApp","Frecuente",3,3,3),
    (3,"Diego","Fernández","097345678","diego@email.com","@diegoF","Paysandú","Paysandú","UY","2024-06-10","2024-12-15",2,2200,820,1100,"Facebook","Ocasional",2,2,2),
    (4,"Ana","Martínez","096456789","ana@email.com","@ana_mtz","Montevideo","Montevideo","UY","2024-08-05","2024-09-10",1,950,350,950,"WooCommerce","Dormido",1,1,2),
    (5,"Carlos","López","095567890","carlos@email.com","@carlos_loop","Rivera","Rivera","UY","2024-02-14","2025-05-20",6,9200,3600,1533,"Instagram","VIP",4,4,5),
    (6,"Sofía","Pérez","094678901","sofia@email.com","@sofia_p","Maldonado","Maldonado","UY","2024-11-01","2025-05-15",4,5600,2100,1400,"WhatsApp","Frecuente",4,3,4),
    (7,"Pablo","González","093789012","pablo@email.com","@pablo_g","Canelones","Canelones","UY","2025-01-10","2025-05-28",2,1800,680,900,"WooCommerce","Ocasional",4,2,2),
    (8,"María","Suárez","092890123","maria@email.com","@maria_s","Tacuarembó","Tacuarembó","UY","2024-04-22","2024-07-30",2,1600,600,800,"Facebook","Dormido",2,2,2),
]

seg_colors = {"VIP":"1B5E20","Frecuente":"0D47A1","Ocasional":"E65100","Dormido":"B71C1C"}
seg_bg = {"VIP":"C8E6C9","Frecuente":"BBDEFB","Ocasional":"FFE0B2","Dormido":"FFCDD2"}

for r_idx, rd in enumerate(sample_clients, 3):
    bg = C_ROW_ALT if r_idx % 2 == 0 else C_ROW_EVEN
    for c_idx, val in enumerate(rd, 1):
        cell = ws3.cell(row=r_idx, column=c_idx, value=val)
        cell.border = border_thin()
        cell.font = font(size=9)
        if c_idx in [10, 11]:
            cell.number_format = "DD/MM/YYYY"
            cell.alignment = align("center", "center")
            cell.fill = fill(bg)
        elif c_idx in [13, 14, 15]:
            cell.number_format = "#,##0"
            cell.alignment = align("right", "center")
            cell.fill = fill(bg)
        elif c_idx == 17:
            seg = val
            cell.fill = fill(seg_bg.get(seg, bg))
            cell.font = font(bold=True, size=9, color=seg_colors.get(seg, C_BLACK))
            cell.alignment = align("center", "center")
        elif c_idx in [18, 19, 20]:
            cell.alignment = align("center", "center")
            cell.fill = fill(bg)
            cell.number_format = "0"
        else:
            cell.fill = fill(bg)
            cell.alignment = align("left" if c_idx in [2,3,4,5,6,7,8,9,16] else "center", "center")
    set_row_height(ws3, r_idx, 16)

# Validación segmento
dv_seg = DataValidation(type="list", formula1='"VIP,Frecuente,Ocasional,Dormido"', allow_blank=True)
ws3.add_data_validation(dv_seg)
dv_seg.add("Q3:Q100000")

dv_canal = DataValidation(type="list",
    formula1='"Instagram,WhatsApp,Facebook,WooCommerce,Mercado Libre,Evento,Referido,Otro"',
    allow_blank=True)
ws3.add_data_validation(dv_canal)
dv_canal.add("P3:P100000")

add_table(ws3, f"A2:T{2+len(sample_clients)}", "TBL_Clientes", "TableStyleMedium7")
freeze(ws3, "D3")

print("Sheet 03 done")

# ═══════════════════════════════════════════════════════
# SHEET 04 — LEADS
# ═══════════════════════════════════════════════════════
ws4 = wb.create_sheet("04_LEADS")
tab_color(ws4, "F57F17")
ws4.sheet_view.showGridLines = False

ws4.merge_cells("A1:L1")
style_title(ws4, "A1", "🎯  LEADS — GESTIÓN DE CONSULTAS", "F57F17", C_WHITE, 13)
set_row_height(ws4, 1, 28)

headers_lead = [
    ("ID Lead",9),("Fecha",12),("Nombre",16),("Teléfono",14),
    ("Instagram",15),("Canal",13),("Producto Interés",22),("Ciudad",14),
    ("Departamento",14),("Estado",16),("Motivo Pérdida",22),("Notas",28),
]

for c_idx, (h, w) in enumerate(headers_lead, 1):
    cell = ws4.cell(row=2, column=c_idx, value=h)
    cell.fill = fill("F57F17")
    cell.font = font(bold=True, color=C_WHITE, size=9)
    cell.alignment = align("center", "center")
    cell.border = border_thin()
    ws4.column_dimensions[get_column_letter(c_idx)].width = w
set_row_height(ws4, 2, 22)

sample_leads = [
    (1,"2025-05-01","Juan Pérez","091111111","@juan_f1","Instagram","Alerón F1","Montevideo","Montevideo","Ganado","",""),
    (2,"2025-05-03","María López","092222222","@maria_l","WhatsApp","Hoodie F1","Salto","Salto","Presupuestado","","Espera talla L"),
    (3,"2025-05-05","Pedro García","093333333","","Facebook","Llavero F1","Paysandú","Paysandú","Perdido","Precio","Muy caro"),
    (4,"2025-05-08","Ana Rodríguez","094444444","@ana_r","WhatsApp","Mate F1","Canelones","Canelones","En negociación","","Quiere pack x2"),
    (5,"2025-05-10","Luis Fernández","095555555","@luis_f","Instagram","Lámpara LED","Montevideo","Montevideo","Contactado","",""),
    (6,"2025-05-12","Carla Suárez","096666666","","WooCommerce","Remera F1","Rivera","Rivera","Nuevo","",""),
    (7,"2025-05-15","Ricardo Martínez","097777777","@ric_mtz","WhatsApp","Circuito Monaco","Maldonado","Maldonado","Ganado","",""),
    (8,"2025-05-18","Valeria Pérez","098888888","@vale_p","Instagram","Vaso Térmico","Montevideo","Montevideo","Perdido","Tiempo","Necesitaba urgente"),
]

estado_colors = {
    "Nuevo":"E3F2FD","Contactado":"FFF9C4","Presupuestado":"FFE0B2",
    "En negociación":"F3E5F5","Ganado":"C8E6C9","Perdido":"FFCDD2"
}
estado_fonts = {
    "Nuevo":"1565C0","Contactado":"F57F17","Presupuestado":"E65100",
    "En negociación":"6A1B9A","Ganado":"1B5E20","Perdido":"B71C1C"
}

for r_idx, rd in enumerate(sample_leads, 3):
    bg = C_ROW_ALT if r_idx % 2 == 0 else C_ROW_EVEN
    for c_idx, val in enumerate(rd, 1):
        cell = ws4.cell(row=r_idx, column=c_idx, value=val)
        cell.border = border_thin()
        cell.font = font(size=9)
        if c_idx == 2:
            cell.number_format = "DD/MM/YYYY"
            cell.alignment = align("center", "center")
            cell.fill = fill(bg)
        elif c_idx == 10:
            estado = val
            cell.fill = fill(estado_colors.get(estado, bg))
            cell.font = font(bold=True, size=9, color=estado_fonts.get(estado, C_BLACK))
            cell.alignment = align("center", "center")
        else:
            cell.fill = fill(bg)
            cell.alignment = align("left" if c_idx in [3,7,11,12] else "center", "center")
    set_row_height(ws4, r_idx, 16)

# Validaciones
dv_estado_lead = DataValidation(type="list",
    formula1='"Nuevo,Contactado,Presupuestado,En negociación,Ganado,Perdido"',
    allow_blank=True)
ws4.add_data_validation(dv_estado_lead)
dv_estado_lead.add("J3:J100000")

dv_canal_lead = DataValidation(type="list",
    formula1='"Instagram,WhatsApp,Facebook,WooCommerce,Mercado Libre,Evento,Referido,Otro"',
    allow_blank=True)
ws4.add_data_validation(dv_canal_lead)
dv_canal_lead.add("F3:F100000")

dv_motivo = DataValidation(type="list",
    formula1='"Precio,Tiempo,Sin respuesta,No le gustó,Competencia,Otro"',
    allow_blank=True)
ws4.add_data_validation(dv_motivo)
dv_motivo.add("K3:K100000")

add_table(ws4, f"A2:L{2+len(sample_leads)}", "TBL_Leads", "TableStyleMedium3")
freeze(ws4, "C3")

print("Sheet 04 done")

# ═══════════════════════════════════════════════════════
# SHEET 05 — COTIZACIONES
# ═══════════════════════════════════════════════════════
ws5 = wb.create_sheet("05_COTIZACIONES")
tab_color(ws5, "6A1B9A")
ws5.sheet_view.showGridLines = False

ws5.merge_cells("A1:K1")
style_title(ws5, "A1", "💰  COTIZACIONES — PRESUPUESTOS ENVIADOS", "6A1B9A", C_WHITE, 13)
set_row_height(ws5, 1, 28)

headers_cot = [
    ("ID Cotización",13),("Fecha",12),("ID Lead",9),("Nombre Cliente",20),
    ("Producto",25),("Cantidad",10),("Precio Unit.",12),("Monto Total",12),
    ("Estado",14),("Fecha Cierre",13),("Notas",28),
]

for c_idx, (h, w) in enumerate(headers_cot, 1):
    cell = ws5.cell(row=2, column=c_idx, value=h)
    cell.fill = fill("6A1B9A")
    cell.font = font(bold=True, color=C_WHITE, size=9)
    cell.alignment = align("center", "center")
    cell.border = border_thin()
    ws5.column_dimensions[get_column_letter(c_idx)].width = w
set_row_height(ws5, 2, 22)

sample_cot = [
    (1,"2025-05-01",1,"Juan Pérez","Alerón F1 Escala 1:18",1,450,450,"Aceptada","2025-05-03","Pagó completo"),
    (2,"2025-05-03",2,"María López","Hoodie F1 talla L",1,2200,2200,"Pendiente","","Espera confirmación"),
    (3,"2025-05-05",3,"Pedro García","Llavero F1 x5",5,90,450,"Rechazada","2025-05-06","Precio muy alto"),
    (4,"2025-05-08",4,"Ana Rodríguez","Mate F1 + Vaso Térmico",2,1400,2800,"Pendiente","","Pack especial"),
    (5,"2025-05-10",5,"Luis Fernández","Lámpara LED F1",1,1800,1800,"Aceptada","2025-05-12","Envío a domicilio"),
    (6,"2025-05-15",7,"Ricardo Martínez","Circuito Monaco Mini",1,380,380,"Aceptada","2025-05-15","Retira en persona"),
    (7,"2025-05-20",2,"María López","Hoodie F1 talla M",1,2200,2200,"Pendiente","","Cambió la talla"),
]

estado_cot_colors = {"Pendiente":"FFF9C4","Aceptada":"C8E6C9","Rechazada":"FFCDD2"}
estado_cot_fonts = {"Pendiente":"F57F17","Aceptada":"1B5E20","Rechazada":"B71C1C"}

for r_idx, rd in enumerate(sample_cot, 3):
    bg = C_ROW_ALT if r_idx % 2 == 0 else C_ROW_EVEN
    for c_idx, val in enumerate(rd, 1):
        cell = ws5.cell(row=r_idx, column=c_idx, value=val)
        cell.border = border_thin()
        cell.font = font(size=9)
        if c_idx in [2, 10]:
            cell.number_format = "DD/MM/YYYY"
            cell.alignment = align("center", "center")
            cell.fill = fill(bg)
        elif c_idx in [7, 8]:
            cell.number_format = "#,##0"
            cell.alignment = align("right", "center")
            cell.fill = fill(bg)
        elif c_idx == 9:
            cell.fill = fill(estado_cot_colors.get(val, bg))
            cell.font = font(bold=True, size=9, color=estado_cot_fonts.get(val, C_BLACK))
            cell.alignment = align("center", "center")
        else:
            cell.fill = fill(bg)
            cell.alignment = align("left" if c_idx in [4,5,11] else "center", "center")
    set_row_height(ws5, r_idx, 16)

dv_estado_cot = DataValidation(type="list",
    formula1='"Pendiente,Aceptada,Rechazada"', allow_blank=True)
ws5.add_data_validation(dv_estado_cot)
dv_estado_cot.add("I3:I100000")

add_table(ws5, f"A2:K{2+len(sample_cot)}", "TBL_Cotizaciones", "TableStyleMedium9")
freeze(ws5, "D3")

print("Sheet 05 done")

# SAVE partial workbook
wb.save("/home/user/pitlane-fastf1-api/erp_temp_part1.xlsx")
print("PART 1 SAVED — sheets 01-05")
